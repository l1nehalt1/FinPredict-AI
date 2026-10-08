"""Bounded CSV/XLSX parsing and deterministic import IDs. No file is executed."""
import base64
import csv
import io
import json
import uuid
import zipfile
from decimal import Decimal, InvalidOperation
from datetime import date, datetime
from collections import Counter

MAX_BYTES=2*1024*1024
MAX_ROWS=3000
HEADERS={
 'date':['date','дата','күні'], 'description':['description','описание','сипаттама'],
 'amount':['amount','сумма','сома'], 'type':['type','тип','түрі'],
 'category':['category','категория','санат'], 'recurring':['recurring','регулярная','тұрақты'],
 'reference':['reference','id','номер операции','операция id']}
TYPES={'income':'income','доход':'income','кіріс':'income','expense':'expense','расход':'expense','шығыс':'expense'}
CATS={'жильё':'housing','продукты':'food','кафе и рестораны':'cafe','покупки':'shopping','транспорт':'transport',
 'подписки и услуги':'services','здоровье':'health','развлечения':'leisure','зарплата':'salary','подработка':'freelance','другое':'other'}

class StatementError(ValueError):pass

def parse(payload, account_id, categories, history_start, as_of, with_metadata=False):
    filename=payload.get('filename','')
    encoded=payload.get('content','')
    if not isinstance(filename,str) or not isinstance(encoded,str) or len(encoded)>MAX_BYTES*4//3+8:
        raise StatementError('Файл слишком большой. Максимум 2 МБ.')
    try:raw=base64.b64decode(encoded,validate=True)
    except (ValueError,TypeError):raise StatementError('Не удалось прочитать файл выписки')
    if len(raw)>MAX_BYTES:raise StatementError('Файл слишком большой. Максимум 2 МБ.')
    if filename.lower().endswith('.pdf'):
        from kaspi import parse_pdf, KaspiError
        try:records,metadata=parse_pdf(raw,account_id,categories)
        except KaspiError as error:raise StatementError(str(error))
        return (records,metadata) if with_metadata else records
    if filename.lower().endswith('.csv'):
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:text=raw.decode('cp1251')
            except UnicodeDecodeError:raise StatementError('Не удалось прочитать файл выписки')
        try:
            dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t')
            rows=list(csv.reader(io.StringIO(text),dialect))
        except csv.Error:raise StatementError('Не удалось прочитать файл выписки')
    elif filename.lower().endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist())>2000 or sum(f.file_size for f in archive.infolist())>20*1024*1024:
                    raise StatementError('Файл слишком большой. Максимум 2 МБ.')
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
            try:
                sheet=book.worksheets[0]
                if sheet.max_column and sheet.max_column>20:raise StatementError('В выписке должно быть не более 20 столбцов')
                rows=[]
                for row in sheet.iter_rows(values_only=True):
                    rows.append(list(row))
                    if len(rows)>MAX_ROWS+1:raise StatementError('Максимум 3000 операций в одном файле')
            finally:book.close()
        except StatementError:raise
        except Exception:raise StatementError('Не удалось прочитать файл выписки')
    else:raise StatementError('Поддерживаются CSV, XLSX и PDF Kaspi Gold')
    rows=[r for r in rows if any(v is not None and str(v).strip() for v in r)]
    if len(rows)<2:raise StatementError('Выписка не содержит операций')
    if len(rows)>MAX_ROWS+1:raise StatementError('Максимум 3000 операций в одном файле')
    headers=[str(v or '').strip().lower() for v in rows[0]]
    if len(headers)>20:raise StatementError('В выписке должно быть не более 20 столбцов')
    if len(headers)!=len(set(headers)):raise StatementError('Названия столбцов не должны повторяться')
    columns={key:next((headers.index(v) for v in values if v in headers),None) for key,values in HEADERS.items()}
    if any(columns[k] is None for k in ['date','description','amount']):
        raise StatementError('Нужны столбцы date, description, amount')
    records=[];seen=Counter()
    for line,row in enumerate(rows[1:],2):
        def val(key):
            index=columns[key]
            return row[index] if index is not None and index<len(row) and row[index] is not None else ''
        try:
            if len(row)>len(headers):raise ValueError()
            d=val('date')
            if isinstance(d,datetime):d=d.date()
            elif not isinstance(d,date):
                d=str(d).strip()
                try:d=date.fromisoformat(d)
                except ValueError:d=datetime.strptime(d,'%d.%m.%Y').date()
            if not history_start<=d<=as_of:raise ValueError()
            description=str(val('description')).strip()
            if not 1<=len(description)<=120 or any(ord(c)<32 for c in description):raise ValueError()
            amount=Decimal(str(val('amount')).replace('\u00a0','').replace(' ','').replace(',','.'))
            if not amount.is_finite() or amount==0 or abs(amount)>100000000 or amount*100!=(amount*100).to_integral_value():raise ValueError()
            kind=str(val('type')).strip().lower()
            kind=TYPES.get(kind) if kind else ('expense' if amount<0 else 'income')
            if not kind:raise ValueError()
            cat=str(val('category')).strip().lower() or ('salary' if kind=='income' else 'other')
            cat=CATS.get(cat,cat)
            if cat not in categories:raise ValueError()
            recurring=str(val('recurring')).strip().lower()
            if recurring not in ['', '0','false','нет','жоқ','1','true','да','иә']:raise ValueError()
            item={'date':d.isoformat(),'description':description,'amount':int(abs(amount)*100),'type':kind,'category':cat,'recurring':recurring in ['1','true','да','иә']}
            reference=str(val('reference')).strip()
            if len(reference)>120:raise ValueError()
            canonical=json.dumps(item,sort_keys=True,ensure_ascii=False)+':'+reference
            seen[canonical]+=1
            item['id']=str(uuid.uuid5(uuid.UUID(account_id),'statement:'+canonical+':'+str(seen[canonical])))
            records.append(item)
        except (ValueError,TypeError,InvalidOperation,OverflowError):
            raise StatementError(f'Строка {line}: проверьте дату, сумму, тип и категорию')
    return (records,None) if with_metadata else records

def safe_csv_cell(value):
    text=str(value)
    return "'"+text if text.lstrip().startswith(('=','+','-','@')) else text
